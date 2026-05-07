import csv
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.print_settings import PrintTitles


DEFAULT_DATASET_NAME = "2025-2026志愿者时长库3月公示"
ROOT = Path(__file__).resolve().parents[1]
DATASET_NAME = os.environ.get("PUBLIC_DATASET_NAME", DEFAULT_DATASET_NAME)
SOURCE_DIR = Path(os.environ.get("PUBLIC_SOURCE_DIR", ROOT / DATASET_NAME))
OUTPUT_DIR = Path(os.environ.get("PUBLIC_OUTPUT_DIR", ROOT / "import-output" / "march-reconcile"))

FORMAL_COLLEGES = [
    "车辆工程学院",
    "机械工程学院",
    "材料科学与工程学院",
    "电气与电子工程学院",
    "计算机科学与工程学院",
    "药学与生物工程学院",
    "化学化工学院",
    "数学科学学院",
    "会计学院",
    "管理学院",
    "创新创业学院",
    "经济金融学院",
    "马克思主义学院",
    "重庆知识产权学院",
    "语言与传播学院",
    "两江国际学院",
    "两江人工智能学院",
    "物理与新能源学院",
    "应用技术学院",
]

ALIASES = {
    "车辆学院": "车辆工程学院",
    "车辆": "车辆工程学院",
    "机械学院": "机械工程学院",
    "机械": "机械工程学院",
    "材料学院": "材料科学与工程学院",
    "材科学院": "材料科学与工程学院",
    "材料科学学院": "材料科学与工程学院",
    "材料": "材料科学与工程学院",
    "电气学院": "电气与电子工程学院",
    "电子与电气学院": "电气与电子工程学院",
    "电气电子学院": "电气与电子工程学院",
    "电气": "电气与电子工程学院",
    "计算机学院": "计算机科学与工程学院",
    "计算机": "计算机科学与工程学院",
    "药生学院": "药学与生物工程学院",
    "药学与生物学院": "药学与生物工程学院",
    "药生": "药学与生物工程学院",
    "化工学院": "化学化工学院",
    "化学学院": "化学化工学院",
    "数学学院": "数学科学学院",
    "数学": "数学科学学院",
    "会计": "会计学院",
    "管理": "管理学院",
    "创新创业": "创新创业学院",
    "经金学院": "经济金融学院",
    "经金": "经济金融学院",
    "马院": "马克思主义学院",
    "马克思学院": "马克思主义学院",
    "知产学院": "重庆知识产权学院",
    "知识产权学院": "重庆知识产权学院",
    "知产": "重庆知识产权学院",
    "语传学院": "语言与传播学院",
    "语传": "语言与传播学院",
    "国际学院": "两江国际学院",
    "两江国际": "两江国际学院",
    "人工智能学院": "两江人工智能学院",
    "人智学院": "两江人工智能学院",
    "人智": "两江人工智能学院",
    "物新学院": "物理与新能源学院",
    "物新": "物理与新能源学院",
    "物理新能源学院": "物理与新能源学院",
    "物理与系能源学院": "物理与新能源学院",
    "应用技术学校": "应用技术学院",
    "应用技术": "应用技术学院",
}

ALIAS_BY_KEY = {}

_PRINT_TITLES_FROM_STRING = PrintTitles.from_string


def _safe_print_titles_from_string(value):
    try:
        return _PRINT_TITLES_FROM_STRING(value)
    except ValueError:
        class EmptyPrintTitles:
            rows = None
            cols = None

        return EmptyPrintTitles()


PrintTitles.from_string = staticmethod(_safe_print_titles_from_string)


def college_key(value):
    return re.sub(r"[。．.，,、；;：:]", "", re.sub(r"\s+", "", str(value or "").replace("\u00a0", " "))).strip()


for formal in FORMAL_COLLEGES:
    ALIAS_BY_KEY[college_key(formal)] = formal
for alias, formal in ALIASES.items():
    ALIAS_BY_KEY[college_key(alias)] = formal


def normalize_college(value):
    text = clean_cell(value)
    if not text:
        return ""
    return ALIAS_BY_KEY.get(college_key(text), text)


def clean_cell(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ").replace("\r", " ").replace("\n", " ").replace("\t", " ")).strip()


def normalize_header(value):
    return re.sub(r"\s+", "", clean_cell(value))


def truncate(value, max_length):
    text = clean_cell(value)
    return text if len(text) <= max_length else text[: max_length - 1] + "…"


def parse_hours(value):
    text = clean_cell(value).replace(",", "").replace("，", "").replace("小时", "")
    text = re.sub(r"h$", "", text, flags=re.I)
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    num = float(match.group(0))
    if num <= 0:
        return None
    return round(num, 1)


def find_header_row(rows):
    best = None
    for idx, row in enumerate(rows[:12]):
        headers = [normalize_header(x) for x in row]
        score = int(any(h == "姓名" or h.endswith("姓名") for h in headers))
        score += int(any(h == "学号" or "学号" in h for h in headers))
        score += int(any("服务时长" in h or "志愿时长" in h or "总时长" in h for h in headers))
        if score >= 2 and (best is None or score > best[1]):
            best = (idx, score)
    return -1 if best is None else best[0]


def find_columns(header_row):
    columns = {"activityColumns": []}
    for idx, raw in enumerate(header_row):
        header = normalize_header(raw)
        if not header:
            continue
        if "student_name" not in columns and (header == "姓名" or header.endswith("姓名")):
            columns["student_name"] = idx
        elif "college" not in columns and header == "学院":
            columns["college"] = idx
        elif "grade" not in columns and header == "年级":
            columns["grade"] = idx
        elif "student_id" not in columns and (header == "学号" or "学号" in header):
            columns["student_id"] = idx
        elif "phone" not in columns and ("联系方式" in header or header == "电话" or "手机号" in header):
            columns["phone"] = idx
        elif "hours" not in columns and ("服务时长" in header or "志愿时长" in header or "总时长" in header):
            columns["hours"] = idx

        is_activity = (
            ("活动" in header or "补录" in header or "寒暑" in header)
            and "服务时长" not in header
            and "志愿时长" not in header
        )
        if is_activity:
            columns["activityColumns"].append({"index": idx, "header": clean_cell(raw)})
    return columns


def required_columns_present(columns):
    return all(key in columns for key in ("student_name", "student_id", "hours"))


def activity_date_for_period(period):
    header = normalize_header(period)
    match = re.search(r"(20\d{2})-(20\d{2})(上|下)半", header)
    if match:
        start_year = int(match.group(1))
        end_year = int(match.group(2))
        return f"{start_year}-12-31" if match.group(3) == "上" else f"{end_year}-06-30"
    single = re.search(r"(20\d{2})", header)
    if single:
        return f"{single.group(1)}-12-31"
    return "2026-03-31"


def normalize_activity_text(value):
    return clean_cell(value).replace("（", "(").replace("）", ")").replace("；", "、")


def repair_malformed_hour_closings(text):
    return re.sub(
        r"(\d+(?:\.\d+)?)\s*(小时|h|H)\s*(?=(?:[、,，。；;]\s*)?[^()]{1,80}\(\s*\d+(?:\.\d+)?\s*(?:小时|h|H))",
        r"\1\2)",
        text,
    )


def clean_activity_name(value):
    name = value.strip()
    name = re.sub(r"^\d+(?:\.\d+)?\s*(?:小时|h|H)\s*[、,，。；;：:\"]*", "", name)
    name = re.sub(r'^[、,，。；;：:"/\\]+|[、,，。；;：:"/\\]+$', "", name)
    return name.strip()


def add_unresolved_activities(fragment, unresolved):
    normalized = fragment.replace("|", "、").replace("｜", "、")
    for part in re.split(r"[、,，。；;]+", normalized):
        name = clean_activity_name(part)
        if len(name) >= 2 and not re.search(r"[()]", name) and not re.search(r"\d+(?:\.\d+)?\s*(?:小时|h|H)", name):
            unresolved.append({"activity_name": truncate(name, 256)})


def extract_activity_details(value):
    text = repair_malformed_hour_closings(normalize_activity_text(value))
    out = []
    unresolved = []
    last_end = 0
    for match in re.finditer(r"([^()]+?)\s*\(\s*(\d+(?:\.\d+)?)\s*(?:小时|h|H)\s*\)", text):
        add_unresolved_activities(text[last_end:match.start()], unresolved)
        name = clean_activity_name(match.group(1))
        hours = round(float(match.group(2)), 1)
        if name and hours > 0:
            out.append({"activity_name": truncate(name, 256), "hours": hours})
        last_end = match.end()
    add_unresolved_activities(text[last_end:], unresolved)
    return {"entries": out, "unresolved": unresolved}


def extract_activities(value):
    return extract_activity_details(value)["entries"]


def get_cell(row, idx):
    return row[idx] if idx is not None and idx < len(row) else ""


def read_workbook_rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    file_name = path.name
    rows_out = []
    invalid_rows = []
    sheet_summaries = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        rows = [row for row in rows if any(clean_cell(cell) for cell in row)]
        if not rows:
            sheet_summaries.append({"sheet": sheet_name, "rows": 0, "imported": 0, "skipped": 0, "reason": "empty"})
            continue
        header_idx = find_header_row(rows)
        if header_idx < 0:
            sheet_summaries.append({"sheet": sheet_name, "rows": len(rows), "imported": 0, "skipped": len(rows), "reason": "header_not_found"})
            continue
        columns = find_columns(rows[header_idx])
        if not required_columns_present(columns):
            sheet_summaries.append({"sheet": sheet_name, "rows": len(rows), "imported": 0, "skipped": len(rows), "reason": "required_columns_missing"})
            continue

        imported = 0
        skipped = 0
        for row_idx in range(header_idx + 1, len(rows)):
            row = rows[row_idx]
            student_name = clean_cell(get_cell(row, columns.get("student_name")))
            student_id = clean_cell(get_cell(row, columns.get("student_id")))
            reported_hours = parse_hours(get_cell(row, columns.get("hours")))
            if not any(clean_cell(cell) for cell in row):
                skipped += 1
                continue
            if not student_name or not student_id or reported_hours is None:
                reason = []
                if not student_name:
                    reason.append("missing_student_name")
                if not student_id:
                    reason.append("missing_student_id")
                if reported_hours is None:
                    reason.append("missing_or_invalid_hours")
                invalid_rows.append({"file": file_name, "sheet": sheet_name, "row_number": row_idx + 1, "reason": ",".join(reason), "row": [clean_cell(x) for x in row]})
                skipped += 1
                continue

            entries = []
            unresolved_entries = []
            detail_length = 0
            for activity_col in columns["activityColumns"]:
                raw_text = clean_cell(get_cell(row, activity_col["index"]))
                detail_length += len(raw_text)
                details = extract_activity_details(raw_text)
                for activity in details["entries"]:
                    entries.append({
                        **activity,
                        "activity_period": activity_col["header"],
                        "activity_date": activity_date_for_period(activity_col["header"]),
                        "source_column": activity_col["header"],
                    })
                for activity in details["unresolved"]:
                    unresolved_entries.append({
                        **activity,
                        "activity_period": activity_col["header"],
                        "activity_date": activity_date_for_period(activity_col["header"]),
                        "source_column": activity_col["header"],
                    })

            rows_out.append({
                "student_id": student_id,
                "student_name": student_name,
                "college": normalize_college(get_cell(row, columns.get("college"))),
                "grade": clean_cell(get_cell(row, columns.get("grade"))),
                "phone": truncate(clean_cell(get_cell(row, columns.get("phone"))), 20),
                "source_file": file_name,
                "source_sheet": sheet_name,
                "source_row": row_idx + 1,
                "source_total_hours": reported_hours,
                "parsed_hours": round(sum(entry["hours"] for entry in entries), 1),
                "detail_length": detail_length,
                "entries": entries,
                "unresolved_entries": unresolved_entries,
            })
            imported += 1
        sheet_summaries.append({"sheet": sheet_name, "rows": len(rows), "header_row": header_idx + 1, "imported": imported, "skipped": skipped})
    return rows_out, invalid_rows, sheet_summaries


def choose_deduped_row(existing, candidate):
    if existing is None:
        return candidate
    if candidate["source_total_hours"] > existing["source_total_hours"]:
        return candidate
    if candidate["source_total_hours"] == existing["source_total_hours"] and candidate["detail_length"] > existing["detail_length"]:
        return candidate
    return existing


def base_record(row, overrides):
    out = {
        "student_id": row["student_id"],
        "student_name": row["student_name"],
        "college": row["college"],
        "grade": row["grade"],
        "phone": row["phone"],
        "source_file": row["source_file"],
        "source_sheet": row["source_sheet"],
        "source_row": row["source_row"],
        "source_total_hours": row["source_total_hours"],
    }
    out.update(overrides)
    return out


def build_import_records(row, audit):
    diff = round(row["source_total_hours"] - row["parsed_hours"], 1)
    if len(row["entries"]) == 0:
        audit["summary_records"] += 1
        audit["no_activity_rows"] += 1
        return [base_record(row, {
            "activity_name": "源表累计时长（未列明活动明细）",
            "activity_period": "源表累计",
            "activity_date": "2026-03-31",
            "hours": row["source_total_hours"],
            "import_detail_status": "summary_no_detail",
        })]
    if diff < -0.05:
        audit["summary_records"] += 1
        audit["conflict_rows"] += 1
        audit["conflict_over_hours"] = round(audit["conflict_over_hours"] + abs(diff), 1)
        audit["conflicts"].append({
            "student_id": row["student_id"],
            "student_name": row["student_name"],
            "source_file": row["source_file"],
            "source_sheet": row["source_sheet"],
            "source_row": row["source_row"],
            "source_total_hours": row["source_total_hours"],
            "parsed_hours": row["parsed_hours"],
            "diff": diff,
            "entries": row["entries"][:20],
        })
        return [base_record(row, {
            "activity_name": "源表累计时长（活动明细合计超过总时长，待人工复核）",
            "activity_period": "源表累计",
            "activity_date": "2026-03-31",
            "hours": row["source_total_hours"],
            "import_detail_status": "summary_conflict",
        })]

    records = [
        base_record(row, {
            "activity_name": entry["activity_name"],
            "activity_period": entry["activity_period"],
            "activity_date": entry["activity_date"],
            "hours": entry["hours"],
            "import_detail_status": "parsed",
        })
        for entry in row["entries"]
    ]
    audit["parsed_activity_records"] += len(records)
    if diff > 0.05:
        audit["supplement_records"] += 1
        audit["supplement_hours"] = round(audit["supplement_hours"] + diff, 1)
        unresolved_activity = row.get("unresolved_entries", [None])[0] if len(row.get("unresolved_entries", [])) == 1 else None
        records.append(base_record(row, {
            "activity_name": unresolved_activity["activity_name"] if unresolved_activity else "源表累计未拆分时长",
            "activity_period": unresolved_activity["activity_period"] if unresolved_activity else "源表累计",
            "activity_date": unresolved_activity["activity_date"] if unresolved_activity else "2026-03-31",
            "hours": diff,
            "import_detail_status": "supplement",
        }))
    return records


LEGACY_PLACEHOLDER_PATTERN = re.compile(r"(源表|未拆分|拆分时长|拆分明细|已拆分|累计未拆分)")


def is_legacy_placeholder(value):
    return bool(LEGACY_PLACEHOLDER_PATTERN.search(clean_cell(value)))


def public_activity_name(value):
    text = clean_cell(value)
    if not text or is_legacy_placeholder(text):
        return "累计时长记录"
    return text


def public_activity_period(value):
    text = clean_cell(value)
    if not text:
        return ""
    if is_legacy_placeholder(text):
        return "累计记录"
    return text


def one_decimal(value):
    if value is None or value == "":
        return ""
    return f"{float(value):.1f}"


CANONICAL_FIELDS = [
    "student_id",
    "student_name",
    "college",
    "grade",
    "phone",
    "source_file",
    "source_sheet",
    "source_row",
    "source_total_hours",
    "activity_name",
    "activity_period",
    "activity_date",
    "hours",
    "import_detail_status",
]


def canonical_value(record, field):
    if field in ("source_total_hours", "hours"):
        return one_decimal(record.get(field))
    value = record.get(field)
    if value is None:
        return ""
    return str(value)


def canonical_line(record):
    return "\t".join(canonical_value(record, field) for field in CANONICAL_FIELDS)


def records_hash(records):
    lines = sorted(canonical_line(record) for record in records)
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def activity_group_key(record):
    return (
        public_activity_name(record["activity_name"]),
        public_activity_period(record["activity_period"]),
        record["activity_date"],
    )


def summarize_activities(records):
    grouped = {}
    for record in records:
        key = activity_group_key(record)
        item = grouped.setdefault(key, {
            "activity_name": key[0],
            "activity_period": key[1],
            "activity_date": key[2],
            "record_count": 0,
            "student_ids": set(),
            "total_hours": 0.0,
            "hours": [],
            "colleges": Counter(),
            "source_files": Counter(),
            "detail_status": Counter(),
            "has_summary_placeholder": False,
        })
        item["record_count"] += 1
        item["student_ids"].add(record["student_id"])
        item["total_hours"] = round(item["total_hours"] + float(record["hours"]), 1)
        item["hours"].append(float(record["hours"]))
        if record.get("college"):
            item["colleges"][record["college"]] += 1
        if record.get("source_file"):
            item["source_files"][record["source_file"]] += 1
        if record.get("import_detail_status"):
            item["detail_status"][record["import_detail_status"]] += 1
        item["has_summary_placeholder"] = item["has_summary_placeholder"] or is_legacy_placeholder(record["activity_name"]) or is_legacy_placeholder(record["activity_period"])

    summary = []
    for item in grouped.values():
        hours = item.pop("hours")
        colleges = item.pop("colleges")
        source_files = item.pop("source_files")
        detail_status = item.pop("detail_status")
        student_ids = item.pop("student_ids")
        item["student_count"] = len(student_ids)
        item["min_hours"] = round(min(hours), 1)
        item["max_hours"] = round(max(hours), 1)
        item["avg_hours"] = round(sum(hours) / len(hours), 2)
        item["college_count"] = len(colleges)
        item["top_colleges"] = "、".join(name for name, _ in colleges.most_common(8))
        item["source_file_count"] = len(source_files)
        item["source_files"] = "、".join(name for name, _ in source_files.most_common(5))
        item["detail_status_counts"] = dict(detail_status)
        item["is_summary_activity"] = item.pop("has_summary_placeholder")
        item["import_description"] = (
            f"来源：{DATASET_NAME}；时间说明：{item['activity_period'] or '未注明'}；"
            f"记录数：{item['record_count']}；学生数：{item['student_count']}；"
            f"累计时长：{item['total_hours']:.1f} 小时；"
            f"单条时长范围：{item['min_hours']:.1f}-{item['max_hours']:.1f} 小时，平均 {item['avg_hours']:.2f} 小时。"
            "同一活动不同学生可对应不同服务时长，活动表不保存固定活动时长。"
        )
        item["start_time"] = item["activity_date"] + "T00:00:00+08:00"
        item["end_time"] = item["activity_date"] + "T23:59:59+08:00"
        item["org_name"] = item["top_colleges"].split("、")[0] if item["college_count"] == 1 and item["top_colleges"] else "多学院/导入汇总"
        summary.append(item)

    summary.sort(key=lambda x: (x["activity_name"], x["activity_period"], x["activity_date"]))
    return summary


def write_csv(path, rows, fields):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not SOURCE_DIR.exists():
        raise SystemExit(f"source dir not found: {SOURCE_DIR}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    files = sorted([p for p in SOURCE_DIR.glob("*.xlsx") if not p.name.startswith("~$")], key=lambda p: p.name)
    raw_rows = []
    invalid_rows = []
    file_summaries = []
    for file in files:
        rows, invalid, sheets = read_workbook_rows(file)
        raw_rows.extend(rows)
        invalid_rows.extend(invalid)
        file_summaries.append({"file": file.name, "imported_raw": len(rows), "invalid": len(invalid), "sheets": sheets})

    groups = defaultdict(list)
    for row in raw_rows:
        groups[(row["student_id"], row["student_name"])].append(row)

    deduped_rows = []
    duplicate_groups = []
    for key, group in groups.items():
        chosen = None
        for row in group:
            chosen = choose_deduped_row(chosen, row)
        deduped_rows.append(chosen)
        if len(group) > 1:
            duplicate_groups.append({
                "key": f"{key[0]} / {key[1]}",
                "count": len(group),
                "chosen": {
                    "student_id": chosen["student_id"],
                    "student_name": chosen["student_name"],
                    "source_total_hours": chosen["source_total_hours"],
                    "parsed_hours": chosen["parsed_hours"],
                    "source_file": chosen["source_file"],
                    "source_sheet": chosen["source_sheet"],
                    "source_row": chosen["source_row"],
                },
            })

    deduped_rows.sort(key=lambda r: (r["source_file"], r["source_sheet"], r["source_row"]))
    audit = {
        "parsed_activity_records": 0,
        "supplement_records": 0,
        "supplement_hours": 0.0,
        "summary_records": 0,
        "no_activity_rows": 0,
        "conflict_rows": 0,
        "conflict_over_hours": 0.0,
        "conflicts": [],
    }
    records = []
    for row in deduped_rows:
        records.extend(build_import_records(row, audit))

    activity_summary = summarize_activities(records)
    source_total_hours = round(sum(row["source_total_hours"] for row in deduped_rows), 1)
    output_total_hours = round(sum(row["hours"] for row in records), 1)
    detail_hash = records_hash(records)
    student_totals_hash = hashlib.sha256(
        "\n".join(
            sorted(
                f"{row['student_id']}\t{row['student_name']}\t{one_decimal(row['source_total_hours'])}"
                for row in deduped_rows
            )
        ).encode("utf-8")
    ).hexdigest()

    meta = {
        "dataset_name": DATASET_NAME,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "files": len(files),
        "raw_student_rows": len(raw_rows),
        "final_students": len(deduped_rows),
        "final_records": len(records),
        "invalid_rows": len(invalid_rows),
        "duplicate_groups": len(duplicate_groups),
        "duplicate_rows_removed": len(raw_rows) - len(deduped_rows),
        "source_total_hours": source_total_hours,
        "output_total_hours": output_total_hours,
        "total_hours_match": source_total_hours == output_total_hours,
        "detail_hash_sha256": detail_hash,
        "student_totals_hash_sha256": student_totals_hash,
        "activity_groups": len(activity_summary),
        **{k: v for k, v in audit.items() if k != "conflicts"},
    }

    (OUTPUT_DIR / "expected-records.json").write_text(json.dumps({"meta": meta, "records": records}, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "source-audit.json").write_text(json.dumps({
        "meta": meta,
        "files": file_summaries,
        "invalid_rows": invalid_rows,
        "duplicate_groups": duplicate_groups,
        "conflicts": audit["conflicts"],
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "activity-summary.json").write_text(json.dumps({"meta": meta, "activities": activity_summary}, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "activities-for-import.json").write_text(json.dumps({
        "meta": meta,
        "activities": [
            {
                "name": row["activity_name"],
                "start_time": row["start_time"],
                "end_time": row["end_time"],
                "location": row["activity_period"],
                "org_name": row["org_name"],
                "description": row["import_description"],
                "source_key": f"{row['activity_name']}\u0000{row['activity_period']}\u0000{row['activity_date']}",
                "summary": row,
            }
            for row in activity_summary
        ],
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    write_csv(
        OUTPUT_DIR / "expected-records.csv",
        records,
        CANONICAL_FIELDS,
    )
    write_csv(
        OUTPUT_DIR / "activity-summary.csv",
        activity_summary,
        [
            "activity_name",
            "activity_period",
            "activity_date",
            "record_count",
            "student_count",
            "total_hours",
            "min_hours",
            "max_hours",
            "avg_hours",
            "college_count",
            "top_colleges",
            "source_file_count",
            "source_files",
            "is_summary_activity",
            "org_name",
            "import_description",
        ],
    )

    print(json.dumps({
        "output_dir": str(OUTPUT_DIR),
        "meta": meta,
        "files": {
            "expected_records": str(OUTPUT_DIR / "expected-records.json"),
            "activity_summary": str(OUTPUT_DIR / "activity-summary.json"),
            "activities_for_import": str(OUTPUT_DIR / "activities-for-import.json"),
            "source_audit": str(OUTPUT_DIR / "source-audit.json"),
        },
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
